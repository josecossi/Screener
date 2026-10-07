#!/usr/bin/env python3
"""
Screener automático de acciones infravaloradas (EE. UU. + Europa).

Qué hace:
  1. Construye el universo con los componentes de varios índices (Wikipedia)
     + los tickers que añadas en tickers_extra.txt.
  2. Descarga datos fundamentales y de precio de Yahoo Finance (librería yfinance).
  3. Convierte todo a EUR, aplica filtros de liquidez y calidad y puntúa 0-100:
       Valoración (PER, EV/EBITDA, P/FCF frente a su sector)
       Calidad    (ROE y deuda neta/EBITDA)
       Momentum   (rentabilidad 12 m y revisiones de estimaciones de BPA)
  4. Genera un Excel con el ranking en la carpeta resultados/ y, opcionalmente,
     lo envía por correo.

AVISOS:
  - yfinance NO es una API oficial de Yahoo: puede fallar, cambiar o dar datos
    erróneos. Verifica siempre las cifras de las candidatas en sus cuentas anuales.
  - Herramienta de análisis, no recomendación de inversión.

Uso:
  python screener.py                  # ejecución completa
  python screener.py --max 50         # prueba rápida con 50 tickers
  python screener.py --no-email       # no enviar correo aunque esté configurado
"""
from __future__ import annotations

import argparse
import logging
import os
import smtplib
import sys
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from email.message import EmailMessage
from io import StringIO
from pathlib import Path

import numpy as np
import pandas as pd
import requests

# =============================================================================
# CONFIGURACIÓN (edita aquí)
# =============================================================================
CONFIG = {
    # Índices que forman el universo: (nombre, URL Wikipedia, sufijo Yahoo)
    # Un sufijo vacío = EE. UU. Puedes comentar los que no quieras.
    "indices": [
        ("S&P 500", "https://en.wikipedia.org/wiki/List_of_S%26P_500_companies", ""),
        ("S&P 400 MidCap", "https://en.wikipedia.org/wiki/List_of_S%26P_400_companies", ""),
        ("IBEX 35", "https://en.wikipedia.org/wiki/IBEX_35", ".MC"),
        ("DAX", "https://en.wikipedia.org/wiki/DAX", ".DE"),
        ("MDAX", "https://en.wikipedia.org/wiki/MDAX", ".DE"),
        ("CAC 40", "https://en.wikipedia.org/wiki/CAC_40", ".PA"),
        ("FTSE MIB", "https://en.wikipedia.org/wiki/FTSE_MIB", ".MI"),
        ("AEX", "https://en.wikipedia.org/wiki/AEX_index", ".AS"),
        ("FTSE 100", "https://en.wikipedia.org/wiki/FTSE_100_Index", ".L"),
    ],
    # Filtros eliminatorios (importes en millones de EUR)
    "cap_min_meur": 1000,          # capitalización mínima
    "vol_min_meur": 1.0,           # volumen medio diario negociado mínimo
    "deuda_ebitda_max": 3.0,       # deuda neta / EBITDA máxima
    "roe_min": 0.10,               # ROE mínimo (proxy de ROIC; yfinance no da ROIC)
    "excluir_financieras": True,   # bancos/aseguradoras: EV/EBITDA no aplica
    # Comparación sectorial
    "min_empresas_sector": 3,
    # Pesos (deben sumar 1)
    "peso_valoracion": 0.40,
    "peso_calidad": 0.30,
    "peso_momentum": 0.30,
    # Salida
    "top_n": 25,
    "workers": 2,                  # descargas en paralelo (más = riesgo de bloqueo)
    "pausa_seg": 0.6,              # pausa entre peticiones por hilo
}

BASE = Path(__file__).resolve().parent
OUT_DIR = BASE / "resultados"
CACHE_DIR = BASE / "cache"
EXTRA_FILE = BASE / "tickers_extra.txt"
UA = {"User-Agent": "Mozilla/5.0 (screener personal; uso no comercial)"}

log = logging.getLogger("screener")


# =============================================================================
# 1. UNIVERSO
# =============================================================================
TICKER_COLS = {"symbol", "ticker", "ticker symbol", "code", "epic"}
SUFIJOS_YAHOO = {".DE", ".F", ".PA", ".MC", ".L", ".AS", ".MI", ".SW", ".BR", ".LS", ".HE",
                 ".CO", ".ST", ".OL", ".VI", ".IR", ".TO", ".MX"}


def _normaliza(tk: str, sufijo: str) -> str | None:
    tk = str(tk).strip().upper().replace(" ", "")
    if not tk or tk == "NAN" or len(tk) > 15:
        return None
    # Si ya trae sufijo de Yahoo (p. ej. AIR.PA dentro del DAX), se respeta
    if "." in tk and "." + tk.rsplit(".", 1)[1] in SUFIJOS_YAHOO:
        return tk
    # Yahoo usa '-' en clases de acciones (BRK.B -> BRK-B, BT.A -> BT-A.L)
    return tk.replace(".", "-") + sufijo


def tickers_de_indice(nombre: str, url: str, sufijo: str) -> list[str]:
    html = requests.get(url, headers=UA, timeout=30).text
    for t in pd.read_html(StringIO(html)):
        cols = {str(c).strip().lower(): c for c in t.columns}
        col = next((cols[c] for c in cols if c in TICKER_COLS), None)
        if col is not None and len(t) >= 20:
            tks = [x for x in (_normaliza(v, sufijo) for v in t[col]) if x]
            log.info("  %-15s %4d tickers", nombre, len(tks))
            return tks
    raise ValueError("no se encontró una tabla con columna de ticker")


def construir_universo() -> pd.DataFrame:
    filas = []
    for nombre, url, sufijo in CONFIG["indices"]:
        try:
            filas += [(t, nombre) for t in tickers_de_indice(nombre, url, sufijo)]
        except Exception as e:  # la estructura de Wikipedia cambia a veces
            log.warning("  %-15s FALLÓ (%s). Añade sus tickers a mano en tickers_extra.txt", nombre, e)
    if EXTRA_FILE.exists():
        extra = [l.split("#")[0].strip().upper() for l in EXTRA_FILE.read_text(encoding="utf-8").splitlines()]
        extra = [t for t in extra if t]
        filas += [(t, "Manual") for t in extra]
        log.info("  %-15s %4d tickers", "Manual", len(extra))
    u = pd.DataFrame(filas, columns=["ticker", "indice"])
    u = u.groupby("ticker", as_index=False)["indice"].agg(lambda s: ", ".join(sorted(set(s))))
    log.info("Universo total (sin duplicados): %d", len(u))
    return u


# =============================================================================
# 2. DESCARGA
# =============================================================================
_fx_cache: dict[str, float] = {}


def eur_por_unidad(divisa: str | None) -> float | None:
    """Cuántos EUR vale 1 unidad de la divisa. GBp/GBX = peniques."""
    import yfinance as yf

    if not divisa:
        return None
    if divisa in ("GBp", "GBX"):
        g = eur_por_unidad("GBP")
        return g / 100 if g else None
    d = divisa.upper()
    if d == "EUR":
        return 1.0
    if d not in _fx_cache:
        try:
            h = yf.Ticker(f"EUR{d}=X").history(period="5d")
            _fx_cache[d] = 1.0 / float(h["Close"].dropna().iloc[-1])
        except Exception:
            _fx_cache[d] = float("nan")
    v = _fx_cache[d]
    return None if v != v else v


def _num(x):
    try:
        x = float(x)
        return x if np.isfinite(x) else None
    except (TypeError, ValueError):
        return None


def _revision(t) -> str | None:
    """Sube / Igual / Baja según revisiones de BPA de los últimos 30 días."""
    try:
        r = t.eps_revisions
        if r is None or r.empty:
            return None
        fila = r.loc["+1y"] if "+1y" in r.index else r.loc["0y"]
        neto = (fila.get("upLast30days", 0) or 0) - (fila.get("downLast30days", 0) or 0)
        return "Sube" if neto > 0 else "Baja" if neto < 0 else "Igual"
    except Exception:
        return None


def descargar_uno(ticker: str) -> dict:
    import yfinance as yf

    for intento in range(3):
        try:
            time.sleep(CONFIG["pausa_seg"])
            t = yf.Ticker(ticker)
            i = t.info or {}
            if not i.get("quoteType"):
                return {"ticker": ticker, "error": "sin datos"}
            precio = _num(i.get("currentPrice") or i.get("regularMarketPrice"))
            acciones = _num(i.get("sharesOutstanding"))
            fx_cot = eur_por_unidad(i.get("currency"))
            fx_fin = eur_por_unidad(i.get("financialCurrency") or i.get("currency"))
            if fx_cot is None or fx_fin is None or precio is None:
                return {"ticker": ticker, "error": "sin precio o divisa"}

            cap = precio * acciones * fx_cot if acciones else (_num(i.get("marketCap")) or 0) * fx_cot
            vol = _num(i.get("averageDailyVolume3Month") or i.get("averageVolume"))
            f = lambda k: (_num(i.get(k)) * fx_fin) if _num(i.get(k)) is not None else None
            deuda, caja, ebitda = f("totalDebt"), f("totalCash"), f("ebitda")
            fcf, beneficio = f("freeCashflow"), f("netIncomeToCommon")
            deuda_neta = (deuda or 0) - (caja or 0) if deuda is not None or caja is not None else None

            return {
                "ticker": ticker,
                "empresa": i.get("shortName") or i.get("longName"),
                "pais": i.get("country"),
                "divisa": i.get("currency"),
                "sector": i.get("sector") or "Desconocido",
                "industria": i.get("industry"),
                "precio": precio,
                "cap_meur": cap / 1e6 if cap else None,
                "vol_meur": vol * precio * fx_cot / 1e6 if vol else None,
                "per": cap / beneficio if beneficio and beneficio > 0 else None,
                "ev_ebitda": (cap + (deuda_neta or 0)) / ebitda if ebitda and ebitda > 0 else None,
                "p_fcf": cap / fcf if fcf and fcf > 0 else None,
                "fcf_meur": fcf / 1e6 if fcf is not None else None,
                "roe": _num(i.get("returnOnEquity")),
                "deuda_neta_ebitda": deuda_neta / ebitda if ebitda and ebitda > 0 and deuda_neta is not None else None,
                "rent_12m": _num(i.get("52WeekChange")),
                "revision": _revision(t),
                "error": None,
            }
        except Exception as e:
            err = str(e)
            time.sleep(2 * (intento + 1))
    return {"ticker": ticker, "error": f"fallo descarga: {err[:80]}"}


def descargar(tickers: list[str]) -> pd.DataFrame:
    hoy = datetime.now().strftime("%Y%m%d")
    cache = CACHE_DIR / f"datos_{hoy}.csv"
    if cache.exists():
        df = pd.read_csv(cache)
        if set(tickers) <= set(df["ticker"]):
            log.info("Usando datos en caché de hoy (%s)", cache.name)
            return df[df["ticker"].isin(tickers)]
    filas, n = [], len(tickers)
    with ThreadPoolExecutor(max_workers=CONFIG["workers"]) as ex:
        futs = {ex.submit(descargar_uno, t): t for t in tickers}
        for k, fu in enumerate(as_completed(futs), 1):
            filas.append(fu.result())
            if k % 50 == 0 or k == n:
                log.info("  descargados %d/%d", k, n)
    # Segunda pasada para los que fallaron por límite de peticiones
    fallidos = [f["ticker"] for f in filas if str(f.get("error") or "").startswith("fallo descarga")]
    if fallidos:
        log.info("  reintentando %d tickers fallidos tras 60 s de pausa", len(fallidos))
        time.sleep(60)
        rep = {}
        for t in fallidos:
            time.sleep(1.0)
            rep[t] = descargar_uno(t)
        filas = [rep.get(f["ticker"], f) for f in filas]
        log.info("  recuperados %d/%d", sum(1 for r in rep.values() if not r.get("error")), len(fallidos))
    campos = ["ticker", "empresa", "pais", "divisa", "sector", "industria", "precio", "cap_meur", "vol_meur",
              "per", "ev_ebitda", "p_fcf", "fcf_meur", "roe", "deuda_neta_ebitda", "rent_12m", "revision", "error"]
    df = pd.DataFrame(filas).reindex(columns=campos)
    CACHE_DIR.mkdir(exist_ok=True)
    df.to_csv(cache, index=False)
    return df


# =============================================================================
# 3. PUNTUACIÓN
# =============================================================================
def _pct(s: pd.Series, mayor_mejor: bool) -> pd.Series:
    """Percentil 0-1 dentro de la serie (empates = media)."""
    v = s.dropna()
    if len(v) <= 1:
        return pd.Series(1.0, index=v.index).reindex(s.index)
    r = (v.rank(method="average") - 1) / (len(v) - 1)
    return (r if mayor_mejor else 1 - r).reindex(s.index)


def _pct_sector(df: pd.DataFrame, col: str) -> pd.Series:
    """Múltiplo más bajo = mejor. Solo valores > 0. Compara en el sector si hay
    suficientes empresas; si no, con todo el universo."""
    s = df[col].where(df[col] > 0)
    universo = _pct(s, mayor_mejor=False)
    out = pd.Series(np.nan, index=df.index)
    for _, idx in df.groupby("sector").groups.items():
        sub = s.loc[idx]
        if sub.notna().sum() >= CONFIG["min_empresas_sector"]:
            out.loc[idx] = _pct(sub, mayor_mejor=False)
        else:
            out.loc[idx] = universo.loc[idx]
    return out


def puntuar(df: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    c = CONFIG
    df = df.copy()
    motivo = pd.Series("", index=df.index)

    def excluye(mask, texto):
        nonlocal motivo
        mask = mask.fillna(True) & (motivo == "")
        motivo = motivo.mask(mask, texto)

    excluye(df["error"].notna(), "Sin datos: " + df["error"].fillna("").astype(str))
    if c["excluir_financieras"]:
        excluye(df["sector"].isin(["Financial Services"]), "Sector financiero (métricas no aplicables)")
    excluye(~(df["cap_meur"] >= c["cap_min_meur"]), "Capitalización insuficiente o desconocida")
    excluye(~(df["vol_meur"] >= c["vol_min_meur"]), "Liquidez insuficiente o desconocida")
    excluye(~(df["fcf_meur"] > 0), "FCF negativo o desconocido")
    excluye(~(df["deuda_neta_ebitda"] <= c["deuda_ebitda_max"]), "Deuda neta/EBITDA alta o desconocida")
    excluye(~(df["roe"] >= c["roe_min"]), "ROE bajo o desconocido")

    excluidas = df.loc[motivo != "", ["ticker", "empresa", "sector"]].assign(motivo=motivo[motivo != ""])
    ok = df.loc[motivo == ""].copy()
    if ok.empty:
        return ok, excluidas

    ok["s_per"] = _pct_sector(ok, "per")
    ok["s_ev_ebitda"] = _pct_sector(ok, "ev_ebitda")
    ok["s_p_fcf"] = _pct_sector(ok, "p_fcf")
    ok["score_valoracion"] = ok[["s_per", "s_ev_ebitda", "s_p_fcf"]].mean(axis=1)

    ok["score_calidad"] = pd.concat(
        [_pct(ok["roe"], True), _pct(ok["deuda_neta_ebitda"], False)], axis=1
    ).mean(axis=1)

    rev = ok["revision"].map({"Sube": 1.0, "Igual": 0.5, "Baja": 0.0})
    ok["score_momentum"] = pd.concat([_pct(ok["rent_12m"], True), rev], axis=1).mean(axis=1)

    comp = ["score_valoracion", "score_calidad", "score_momentum"]
    sin_comp = ok[comp].isna().any(axis=1)
    excluidas = pd.concat([excluidas, ok.loc[sin_comp, ["ticker", "empresa", "sector"]]
                           .assign(motivo="Faltan datos para puntuar")])
    ok = ok.loc[~sin_comp]
    ok["puntuacion"] = (100 * (c["peso_valoracion"] * ok["score_valoracion"]
                               + c["peso_calidad"] * ok["score_calidad"]
                               + c["peso_momentum"] * ok["score_momentum"])).round(1)
    ok = ok.sort_values("puntuacion", ascending=False).reset_index(drop=True)
    ok.insert(0, "ranking", range(1, len(ok) + 1))
    return ok, excluidas


# =============================================================================
# 4. SALIDA
# =============================================================================
COLS_RANKING = {
    "ranking": "Pos.", "ticker": "Ticker", "empresa": "Empresa", "indice": "Índice", "pais": "País",
    "sector": "Sector", "puntuacion": "Puntuación", "score_valoracion": "Valoración",
    "score_calidad": "Calidad", "score_momentum": "Momentum", "precio": "Precio", "divisa": "Divisa",
    "cap_meur": "Cap. (M€)", "per": "PER", "ev_ebitda": "EV/EBITDA", "p_fcf": "P/FCF", "roe": "ROE",
    "deuda_neta_ebitda": "DN/EBITDA", "rent_12m": "Rent. 12m", "revision": "Revisión BPA",
}
FORMATOS = {"Puntuación": "0.0", "Valoración": "0.00", "Calidad": "0.00", "Momentum": "0.00",
            "Precio": "#,##0.00", "Cap. (M€)": "#,##0", "PER": "0.0x", "EV/EBITDA": "0.0x",
            "P/FCF": "0.0x", "ROE": "0.0%", "DN/EBITDA": "0.0x", "Rent. 12m": "0.0%"}


def _formatea(ws, df):
    from openpyxl.styles import Font, PatternFill, Alignment
    hdr = PatternFill("solid", fgColor="1F4E78")
    for j, col in enumerate(df.columns, 1):
        c = ws.cell(row=1, column=j)
        c.font, c.fill = Font(name="Arial", bold=True, color="FFFFFF"), hdr
        c.alignment = Alignment(wrap_text=True, horizontal="center")
        ancho = max(10, min(32, int(df[col].astype(str).str.len().quantile(0.9) if len(df) else 10) + 2))
        ws.column_dimensions[c.column_letter].width = ancho
        fmt = FORMATOS.get(col)
        for i in range(2, len(df) + 2):
            cell = ws.cell(row=i, column=j)
            cell.font = Font(name="Arial")
            if fmt:
                cell.number_format = fmt
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions


def exportar(ok, excluidas, universo, inicio) -> Path:
    OUT_DIR.mkdir(exist_ok=True)
    ruta = OUT_DIR / f"screener_{datetime.now():%Y-%m-%d}.xlsx"
    ok = ok.merge(universo, on="ticker", how="left")
    cols = [k for k in COLS_RANKING if k in ok.columns]
    top = ok[cols].head(CONFIG["top_n"]).rename(columns=COLS_RANKING)
    todas = ok[cols].rename(columns=COLS_RANKING)
    info = pd.DataFrame({
        "Campo": ["Fecha de ejecución", "Duración (min)", "Empresas analizadas", "Superan filtros",
                  "Fuente de datos", "Fuente del universo", "Parámetros", "Avisos"],
        "Valor": [f"{inicio:%Y-%m-%d %H:%M}", round((datetime.now() - inicio).seconds / 60, 1),
                  len(universo), len(ok),
                  "Yahoo Finance vía yfinance (API no oficial; puede contener errores)",
                  "Componentes de índices en Wikipedia + tickers_extra.txt",
                  str({k: v for k, v in CONFIG.items() if k != "indices"}),
                  "No es recomendación de inversión. Verifica cifras en cuentas anuales. "
                  "Comprueba que el valor está disponible en tu bróker. ROE se usa como proxy de ROIC. "
                  "Barata no implica que vaya a subir a corto plazo."],
    })
    with pd.ExcelWriter(ruta, engine="openpyxl") as w:
        for nombre, d in [("Ranking", top), ("Todas", todas), ("Excluidas", excluidas), ("Info", info)]:
            d.to_excel(w, sheet_name=nombre, index=False)
            if nombre in ("Ranking", "Todas"):
                _formatea(w.sheets[nombre], d)
    # CSV para revisión automática y comparación entre días
    ok.to_csv(OUT_DIR / f"ranking_{datetime.now():%Y-%m-%d}.csv", index=False)
    excluidas.to_csv(OUT_DIR / f"excluidas_{datetime.now():%Y-%m-%d}.csv", index=False)
    log.info("Excel generado: %s", ruta)
    return ruta


def enviar_email(ruta: Path, top: pd.DataFrame):
    usuario, clave, dest = (os.environ.get(k) for k in ("SCREENER_EMAIL_USER", "SCREENER_EMAIL_PASS", "SCREENER_EMAIL_TO"))
    if not (usuario and clave and dest):
        log.info("Correo no configurado (variables SCREENER_EMAIL_*); se omite el envío")
        return
    msg = EmailMessage()
    msg["Subject"] = f"Screener acciones {datetime.now():%d/%m/%Y}: top {min(10, len(top))}"
    msg["From"], msg["To"] = usuario, dest
    lineas = [f"{r.ranking:>2}. {r.ticker:<9} {str(r.empresa)[:28]:<28} {r.puntuacion:>5.1f}" for r in top.head(10).itertuples()]
    msg.set_content("Top 10 del screener (detalle en el Excel adjunto):\n\n" + "\n".join(lineas)
                    + "\n\nAnálisis automático, no recomendación de inversión. Verifica los datos antes de operar.")
    msg.add_attachment(ruta.read_bytes(), maintype="application",
                       subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet", filename=ruta.name)
    with smtplib.SMTP_SSL(os.environ.get("SCREENER_SMTP", "smtp.gmail.com"), 465) as s:
        s.login(usuario, clave)
        s.send_message(msg)
    log.info("Correo enviado a %s", dest)


# =============================================================================
def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--max", type=int, help="limitar nº de tickers (pruebas)")
    p.add_argument("--no-email", action="store_true")
    a = p.parse_args()

    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(message)s", datefmt="%H:%M:%S")
    pesos = CONFIG["peso_valoracion"] + CONFIG["peso_calidad"] + CONFIG["peso_momentum"]
    if abs(pesos - 1) > 1e-6:
        sys.exit(f"Los pesos suman {pesos}, deben sumar 1")

    inicio = datetime.now()
    log.info("1/4 Construyendo universo")
    universo = construir_universo()
    if universo.empty:
        sys.exit("Universo vacío: revisa la conexión o tickers_extra.txt")
    if a.max:
        universo = universo.head(a.max)
    log.info("2/4 Descargando datos de %d empresas (puede tardar 15-40 min)", len(universo))
    datos = descargar(universo["ticker"].tolist())
    log.info("3/4 Puntuando")
    ok, excluidas = puntuar(datos)
    log.info("   %d superan filtros, %d excluidas", len(ok), len(excluidas))
    log.info("4/4 Exportando")
    ruta = exportar(ok, excluidas, universo, inicio)
    if not ok.empty:
        print("\nTOP 10\n" + ok.head(10)[["ranking", "ticker", "empresa", "sector", "puntuacion"]].to_string(index=False))
        if not a.no_email:
            enviar_email(ruta, ok)


if __name__ == "__main__":
    main()
